# -*- coding: utf-8 -*-
"""Construye el documento DOCX de OmniSafety AI a partir de Proyecto_OmniSafety_AI.txt.

El archivo fuente emplea marcadores simples que este script interpreta:

    [CUBIERTA] ... [FIN_CUBIERTA]   portada
    [PAGINA]                        salto de página
    [INDICE]                        campo de índice automático
    [H1] / [H2] / [H3]              títulos
    [TABLA] N | Título              tabla (las filas comienzan con "|")
    [FIGURA] archivo.png | Título   ilustración desde docs/figuras/
    [NOTA] texto                    nota al pie de tabla o figura
    [MONO] texto                    bloque de código
    [REF] texto                     referencia con sangría francesa

Salida: Proyecto_OmniSafety_AI.docx
"""

from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Permite ejecutar el script tanto desde la raíz como desde tools/.
RAIZ_PROYECTO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if RAIZ_PROYECTO not in sys.path:
    sys.path.insert(0, RAIZ_PROYECTO)

from src import paths  # noqa: E402

FUENTE = "Times New Roman"
FUENTE_TABLA = "Times New Roman"
FUENTE_MONO = "Consolas"

SRC = paths.FUENTE_DOCUMENTO
DST = paths.SALIDA_DOCUMENTO
DIR_FIGURAS = paths.DOCS_FIGURAS


# ---------------------------------------------------------------------------
# Utilidades de bajo nivel para XML
# ---------------------------------------------------------------------------
def _sombrear(elemento, color: str) -> None:
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:color"), "auto")
    sombreado.set(qn("w:fill"), color)
    elemento.append(sombreado)


def _campo(parrafo, instruccion: str, marcador: str = "") -> None:
    """Inserta un campo de Word (PAGE, TOC, etc.) dentro del párrafo."""
    corrida = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    corrida._r.append(inicio)

    corrida = parrafo.add_run()
    texto_instruccion = OxmlElement("w:instrText")
    texto_instruccion.set(qn("xml:space"), "preserve")
    texto_instruccion.text = instruccion
    corrida._r.append(texto_instruccion)

    corrida = parrafo.add_run()
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    corrida._r.append(separador)

    if marcador:
        corrida = parrafo.add_run(marcador)

    corrida = parrafo.add_run()
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    corrida._r.append(fin)


def _activar_actualizacion_de_campos(documento) -> None:
    ajustes = documento.settings.element
    actualizar = OxmlElement("w:updateFields")
    actualizar.set(qn("w:val"), "true")
    ajustes.append(actualizar)


def _configurar_fuente(corrida, tamano=12, negrita=False, cursiva=False, fuente=FUENTE) -> None:
    corrida.font.name = fuente
    corrida.font.size = Pt(tamano)
    corrida.bold = negrita
    corrida.italic = cursiva
    corrida.font.color.rgb = RGBColor(0, 0, 0)
    rpr = corrida._element.get_or_add_rPr()
    fuentes = rpr.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        rpr.append(fuentes)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(atributo), fuente)


# ---------------------------------------------------------------------------
# Estilos globales del documento
# ---------------------------------------------------------------------------
def configurar_estilos(documento) -> None:
    normal = documento.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rpr = normal.element.get_or_add_rPr()
    fuentes = rpr.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        rpr.append(fuentes)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(atributo), FUENTE)

    formato = normal.paragraph_format
    formato.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    formato.space_after = Pt(0)
    formato.space_before = Pt(0)
    formato.first_line_indent = Inches(0.5)

    especificaciones = {
        "Heading 1": (12, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": (12, True, False, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 3": (12, True, True, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 4": (12, True, True, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for nombre, (tamano, negrita, cursiva, alineacion) in especificaciones.items():
        estilo = documento.styles[nombre]
        estilo.font.name = FUENTE
        estilo.font.size = Pt(tamano)
        estilo.font.bold = negrita
        estilo.font.italic = cursiva
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        formato_estilo = estilo.paragraph_format
        formato_estilo.first_line_indent = Inches(0)
        formato_estilo.alignment = alineacion
        formato_estilo.keep_with_next = True
        formato_estilo.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        if nombre == "Heading 1":
            formato_estilo.space_before = Pt(0)
            formato_estilo.space_after = Pt(12)
        else:
            formato_estilo.space_before = Pt(12)
            formato_estilo.space_after = Pt(6)

    for nombre in ("List Paragraph",):
        estilo = documento.styles[nombre]
        estilo.font.name = FUENTE
        estilo.font.size = Pt(12)


def configurar_pagina(documento) -> None:
    seccion = documento.sections[0]
    seccion.top_margin = Inches(1)
    seccion.bottom_margin = Inches(1)
    seccion.left_margin = Inches(1)
    seccion.right_margin = Inches(1)

    encabezado = seccion.header
    parrafo = encabezado.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    parrafo.paragraph_format.first_line_indent = Inches(0)
    parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _campo(parrafo, "PAGE", "1")
    for corrida in parrafo.runs:
        _configurar_fuente(corrida, tamano=12)


# ---------------------------------------------------------------------------
# Constructores de bloques de contenido
# ---------------------------------------------------------------------------
class Constructor:
    def __init__(self, documento) -> None:
        self.doc = documento
        self.contador_figuras = 0
        self.contador_h1 = 0
        self.ultimo_fue_salto = False
        self.tablas = 0
        self.figuras_insertadas = 0
        self.advertencias: list[str] = []

    # -- bloques básicos ----------------------------------------------------
    def cuerpo(self, texto: str) -> None:
        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        formato = parrafo.paragraph_format
        formato.first_line_indent = Inches(0.5)
        formato.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        formato.space_after = Pt(0)
        _configurar_fuente(parrafo.add_run(texto))
        self.ultimo_fue_salto = False

    def elemento_de_lista(self, texto: str) -> None:
        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        formato = parrafo.paragraph_format
        formato.left_indent = Inches(0.75)
        formato.first_line_indent = Inches(-0.25)
        formato.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        formato.space_after = Pt(0)
        _configurar_fuente(parrafo.add_run(texto))
        self.ultimo_fue_salto = False

    def titulo(self, nivel: int, texto: str) -> None:
        if nivel == 1:
            self.contador_h1 += 1
            if not self.ultimo_fue_salto and self.contador_h1 > 1:
                parrafo = self.doc.add_paragraph()
                parrafo.paragraph_format.page_break_before = True
                parrafo.paragraph_format.first_line_indent = Inches(0)
                parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        parrafo = self.doc.add_heading(level=nivel)
        _configurar_fuente(parrafo.add_run(texto), tamano=12, negrita=True, cursiva=nivel >= 3)
        self.ultimo_fue_salto = False

    def salto_de_pagina(self) -> None:
        self.doc.add_page_break()
        self.ultimo_fue_salto = True

    def indice(self) -> None:
        parrafo = self.doc.add_paragraph()
        parrafo.paragraph_format.first_line_indent = Inches(0)
        parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        _campo(parrafo, 'TOC \\o "1-3" \\h \\z \\u', "Índice generado automáticamente. Pulse F9 para actualizarlo.")
        self.ultimo_fue_salto = False

    def mono(self, texto: str) -> None:
        parrafo = self.doc.add_paragraph()
        formato = parrafo.paragraph_format
        formato.left_indent = Inches(0.35)
        formato.first_line_indent = Inches(0)
        formato.line_spacing_rule = WD_LINE_SPACING.SINGLE
        formato.space_before = Pt(6)
        formato.space_after = Pt(6)
        _configurar_fuente(parrafo.add_run(texto), tamano=10, fuente=FUENTE_MONO)
        _sombrear(parrafo._p.get_or_add_pPr(), "F2F2F2")
        self.ultimo_fue_salto = False

    def referencia(self, texto: str) -> None:
        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        formato = parrafo.paragraph_format
        formato.left_indent = Inches(0.5)
        formato.first_line_indent = Inches(-0.5)
        formato.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        formato.space_after = Pt(0)
        _configurar_fuente(parrafo.add_run(texto))
        self.ultimo_fue_salto = False

    def nota(self, texto: str) -> None:
        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        formato = parrafo.paragraph_format
        formato.first_line_indent = Inches(0)
        formato.line_spacing_rule = WD_LINE_SPACING.SINGLE
        formato.space_before = Pt(4)
        formato.space_after = Pt(12)
        _configurar_fuente(parrafo.add_run("Nota. "), tamano=10, cursiva=True)
        _configurar_fuente(parrafo.add_run(texto), tamano=10)
        self.ultimo_fue_salto = False

    # -- tablas -------------------------------------------------------------
    def tabla(self, numero: str, titulo: str, filas: list[list[str]]) -> None:
        parrafo = self.doc.add_paragraph()
        parrafo.paragraph_format.first_line_indent = Inches(0)
        parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        parrafo.paragraph_format.space_before = Pt(12)
        parrafo.paragraph_format.keep_with_next = True
        _configurar_fuente(parrafo.add_run(f"Tabla {numero}"), tamano=12, negrita=True)

        parrafo = self.doc.add_paragraph()
        parrafo.paragraph_format.first_line_indent = Inches(0)
        parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        parrafo.paragraph_format.keep_with_next = True
        _configurar_fuente(parrafo.add_run(titulo), tamano=12, cursiva=True)

        columnas = max(len(f) for f in filas)
        tabla = self.doc.add_table(rows=0, cols=columnas)
        tabla.style = "Table Grid"
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.autofit = True

        for indice, fila in enumerate(filas):
            celdas = tabla.add_row().cells
            for posicion in range(columnas):
                contenido = fila[posicion] if posicion < len(fila) else ""
                celda = celdas[posicion]
                celda.text = ""
                parrafo_celda = celda.paragraphs[0]
                parrafo_celda.paragraph_format.first_line_indent = Inches(0)
                parrafo_celda.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                parrafo_celda.paragraph_format.space_after = Pt(2)
                parrafo_celda.paragraph_format.space_before = Pt(2)
                parrafo_celda.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _configurar_fuente(
                    parrafo_celda.add_run(contenido),
                    tamano=10,
                    negrita=(indice == 0),
                    fuente=FUENTE_TABLA,
                )
                if indice == 0:
                    _sombrear(celda._tc.get_or_add_tcPr(), "DCE6F1")
        self.tablas += 1
        self.ultimo_fue_salto = False

    # -- figuras ------------------------------------------------------------
    def figura(self, archivo: str, titulo: str) -> None:
        self.contador_figuras += 1
        numero = self.contador_figuras

        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        parrafo.paragraph_format.first_line_indent = Inches(0)
        parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        parrafo.paragraph_format.space_before = Pt(12)
        parrafo.paragraph_format.keep_with_next = True
        _configurar_fuente(parrafo.add_run(f"Figura {numero}"), tamano=12, negrita=True)

        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        parrafo.paragraph_format.first_line_indent = Inches(0)
        parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        parrafo.paragraph_format.keep_with_next = True
        _configurar_fuente(parrafo.add_run(titulo), tamano=12, cursiva=True)

        ruta = os.path.join(DIR_FIGURAS, archivo)
        if os.path.exists(ruta):
            parrafo_imagen = self.doc.add_paragraph()
            parrafo_imagen.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo_imagen.paragraph_format.first_line_indent = Inches(0)
            parrafo_imagen.paragraph_format.space_before = Pt(6)
            parrafo_imagen.paragraph_format.space_after = Pt(4)
            parrafo_imagen.add_run().add_picture(ruta, width=Inches(6.0))
            self.figuras_insertadas += 1
        else:
            self.advertencias.append(f"Figura no encontrada: {ruta}")
        self.ultimo_fue_salto = False

    # -- portada ------------------------------------------------------------
    def cubierta(self, lineas: list[str]) -> None:
        for _ in range(4):
            parrafo = self.doc.add_paragraph()
            parrafo.paragraph_format.first_line_indent = Inches(0)
            parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        for posicion, linea in enumerate(lineas):
            parrafo = self.doc.add_paragraph()
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo.paragraph_format.first_line_indent = Inches(0)
            parrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            if posicion == 0:
                _configurar_fuente(parrafo.add_run(linea), tamano=14, negrita=True)
            else:
                _configurar_fuente(parrafo.add_run(linea), tamano=12)
        self.ultimo_fue_salto = False


# ---------------------------------------------------------------------------
# Análisis del archivo fuente
# ---------------------------------------------------------------------------
def construir_desde_texto(lineas: list[str]) -> tuple[Document, Constructor]:
    documento = Document()
    configurar_estilos(documento)
    configurar_pagina(documento)
    constructor = Constructor(documento)

    indice = 0
    total = len(lineas)
    while indice < total:
        linea = lineas[indice].rstrip("\n")
        despojado = linea.strip()

        if not despojado:
            indice += 1
            continue

        if despojado == "[CUBIERTA]":
            indice += 1
            contenido: list[str] = []
            while indice < total and lineas[indice].strip() != "[FIN_CUBIERTA]":
                if lineas[indice].strip():
                    contenido.append(lineas[indice].strip())
                indice += 1
            constructor.cubierta(contenido)
            indice += 1
            continue

        if despojado == "[FIN_CUBIERTA]":
            indice += 1
            continue

        if despojado == "[PAGINA]":
            constructor.salto_de_pagina()
            indice += 1
            continue

        if despojado == "[INDICE]":
            constructor.indice()
            indice += 1
            continue

        coincidencia = re.match(r"^\[H([123])\]\s*(.+)$", despojado)
        if coincidencia:
            constructor.titulo(int(coincidencia.group(1)), coincidencia.group(2).strip())
            indice += 1
            continue

        coincidencia = re.match(r"^\[REF\]\s*(.+)$", despojado)
        if coincidencia:
            constructor.referencia(coincidencia.group(1).strip())
            indice += 1
            continue

        coincidencia = re.match(r"^\[MONO\]\s*(.+)$", despojado)
        if coincidencia:
            constructor.mono(coincidencia.group(1).strip())
            indice += 1
            continue

        coincidencia = re.match(r"^\[NOTA\]\s*(.+)$", despojado)
        if coincidencia:
            constructor.nota(coincidencia.group(1).strip())
            indice += 1
            continue

        coincidencia = re.match(r"^\[TABLA\]\s*([^|]+)\|\s*(.+)$", despojado)
        if coincidencia:
            numero = coincidencia.group(1).strip()
            titulo = coincidencia.group(2).strip()
            indice += 1
            filas: list[list[str]] = []
            while indice < total:
                candidata = lineas[indice].strip()
                if not candidata.startswith("|"):
                    break
                if re.fullmatch(r"\|[\s\-:|]+\|", candidata):
                    indice += 1
                    continue
                celdas = [c.strip() for c in candidata.strip("|").split("|")]
                filas.append(celdas)
                indice += 1
            constructor.tabla(numero, titulo, filas)
            continue

        coincidencia = re.match(r"^\[FIGURA\]\s*([^|]+)\|\s*(.+)$", despojado)
        if coincidencia:
            constructor.figura(coincidencia.group(1).strip(), coincidencia.group(2).strip())
            indice += 1
            continue

        if re.match(r"^\d+\.\s", despojado):
            constructor.elemento_de_lista(despojado)
            indice += 1
            continue

        if despojado.startswith("- "):
            constructor.elemento_de_lista("• " + despojado[2:].strip())
            indice += 1
            continue

        constructor.cuerpo(despojado)
        indice += 1

    return documento, constructor


def main() -> int:
    if not os.path.exists(SRC):
        print(f"ERROR: no se encontró {SRC}", file=sys.stderr)
        return 1

    with open(SRC, "r", encoding="utf-8") as manejador:
        lineas = manejador.readlines()

    documento, constructor = construir_desde_texto(lineas)
    _activar_actualizacion_de_campos(documento)
    documento.save(DST)

    tamano = os.path.getsize(DST)
    print(f"Documento generado: {DST}")
    print(f"  Párrafos:            {len(documento.paragraphs)}")
    print(f"  Tablas:              {constructor.tablas}")
    print(f"  Figuras insertadas:  {constructor.figuras_insertadas}")
    print(f"  Tamaño del archivo:  {tamano / 1024:.1f} KB")
    if constructor.advertencias:
        print("  Advertencias:")
        for advertencia in constructor.advertencias:
            print(f"    - {advertencia}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
